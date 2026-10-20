"""Helios-style HashiCorp deck kit (python-pptx). Dark theme, Terraform purple primary, product-color accents.

Usage: see example_helios_deck.py. Layouts: title, agenda, section, cards, rows, statement, big, cascade,
scorecard, shoutouts, welcome, close. NO EM DASHES anywhere in slide copy.

Self-contained: needs only python-pptx and Pillow. No image files are required.
- Terraform mark is drawn natively with shapes (no PNG, no network).
- Logo is optional: pass the official lockup file if you have it, otherwise a drawn mark plus "HCP Terraform" text is used.
- Photos are optional: missing photos become initials circles.
- Fonts: uses HashiCorp Sans if installed, otherwise falls back to Arial (the pptx does not embed fonts).

Rules baked in: sentence case labels, one-line bullets, arrows drawn as shapes (glyphs render as dashes),
speaker notes go to a sidecar .md by default (python-pptx notes break Keynote import; opt in with embed_notes=True
for PowerPoint-only decks), theme/master fonts rewritten so no Calibri/Arial/Helvetica leaks in.
Works in PowerPoint (default target) and Keynote.
"""
import glob, os, platform, re, shutil, tempfile, zipfile
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from PIL import Image, ImageDraw, ImageFont

BG, FG, MUTE, LINE, CARD = "000000", "FFFFFF", "B2B6BD", "3B3D45", "0C0C0E"
PURPLE, PURPLE_DEEP = "B87FFF", "7B42BC"                      # text/accents on black, fills/gradient
# Terraform purple, Nomad green, Packer blue, Waypoint teal, Vault yellow, Consul pink (pink/yellow: fills or sparing use)
PAL = ["B87FFF", "00CA8E", "02A8EF", "14C6CB", "FFCF25", "E03875"]


def _have_hashicorp_sans():
    dirs = {"Darwin": ["~/Library/Fonts", "/Library/Fonts"], "Windows": [r"C:\Windows\Fonts", "~\\AppData\\Local\\Microsoft\\Windows\\Fonts"]}.get(platform.system(), ["~/.fonts", "~/.local/share/fonts", "/usr/share/fonts"])
    return any(glob.glob(os.path.join(os.path.expanduser(d), "**", "HashiCorpSans*"), recursive=True) for d in dirs)


if _have_hashicorp_sans():
    HEAD, BODY, MONO = "HashiCorp Sans SemiBold", "HashiCorp Sans", "Menlo"
else:
    HEAD, BODY, MONO = "Arial", "Arial", "Courier New"
    print("helios_kit: HashiCorp Sans not found, falling back to Arial. Run install_hashicorp_sans.py for the real look.")

# Terraform mark polygons in the official 24x24 icon grid (x 2..22, y 0.6..23.4)
MARK = [[(8.893, 4.618), (15.107, 8.215), (15.107, 15.415), (8.893, 11.815)],
        [(15.788, 8.215), (15.788, 15.415), (22, 11.815), (22, 4.618)],
        [(2, 0.6), (2, 7.797), (8.212, 11.397), (8.212, 4.199)],
        [(8.893, 19.8), (15.105, 23.4), (15.105, 16.203), (8.893, 12.603)]]


def rgb(h):
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


class Deck:
    def __init__(self, banner, logo=None, mark=True, photos_dir=None, embed_notes=False):
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = Inches(13.33), Inches(7.5)
        self.banner, self.logo, self.mark, self.photos, self.embed_notes = banner, logo, mark, photos_dir, embed_notes
        self.counters, self.notes_log = [], []

    # ---------- primitives
    def txt(self, s, text, l, t, w, h, size=16, color=FG, font=BODY, align=None, wrap=True):
        tb = s.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h)); tf = tb.text_frame; tf.word_wrap = wrap
        p = tf.paragraphs[0]
        if align: p.alignment = align
        r = p.add_run(); r.text = text; r.font.size = Pt(size); r.font.name = font; r.font.bold = False
        r.font.color.rgb = rgb(color); return tb

    def shape(self, s, kind, l, t, w, h, fill=None, line=None, lw=0.75):
        sh = s.shapes.add_shape(kind, Inches(l), Inches(t), Inches(w), Inches(h))
        if fill is None: sh.fill.background()
        else: sh.fill.solid(); sh.fill.fore_color.rgb = rgb(fill)
        if line is None: sh.line.fill.background()
        else: sh.line.color.rgb = rgb(line); sh.line.width = Pt(lw)
        sh.shadow.inherit = False; return sh

    def rect(self, s, l, t, w, h, fill=None, line=None, lw=0.75):
        return self.shape(s, MSO_SHAPE.RECTANGLE, l, t, w, h, fill, line, lw)

    def draw_mark(self, s, x, y, h, color=PURPLE_DEEP):  # native Terraform mark, h inches tall
        k = h / 22.8
        for pts in MARK:
            P = [(Inches(x + (px - 2) * k), Inches(y + (py - 0.6) * k)) for px, py in pts]
            fb = s.shapes.build_freeform(P[0][0], P[0][1]); fb.add_line_segments(P[1:], close=True)
            sh = fb.convert_to_shape(); sh.fill.solid(); sh.fill.fore_color.rgb = rgb(color); sh.line.fill.background()
            sh.shadow.inherit = False

    def arrow(self, s, x, y, color=MUTE):
        self.shape(s, MSO_SHAPE.RIGHT_ARROW, x, y, 0.32, 0.2, fill=color)

    def panel(self, s, l, t, w, h, bar=None):
        self.rect(s, l, t, w, h, CARD, LINE, 1)
        if bar: self.rect(s, l, t, w, 0.07, fill=bar)

    def tag(self, s, text, l, t):  # outlined mono tag; width scales with text so it never wraps
        w = max(0.9, 0.45 + 0.145 * len(text))
        self.rect(s, l, t, w, 0.38, None, PURPLE, 1)
        self.txt(s, text, l + 0.2, t + 0.06, w - 0.2, 0.3, 12, PURPLE, MONO, wrap=False)

    def new(self, mark=False):
        s = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        s.background.fill.solid(); s.background.fill.fore_color.rgb = rgb(BG)
        b = self.rect(s, 0, 0, 13.33, 0.34); b.fill.gradient(); b.fill.gradient_angle = 0
        b.fill.gradient_stops[0].color.rgb = rgb(PURPLE_DEEP); b.fill.gradient_stops[1].color.rgb = rgb(PURPLE)
        self.txt(s, self.banner, 0, 0.03, 13.33, 0.3, 11, FG, BODY, PP_ALIGN.CENTER)
        if self.logo and os.path.exists(self.logo):  # official lockup if supplied
            s.shapes.add_picture(self.logo, Inches(11.2), Inches(0.5), width=Inches(1.9), height=Inches(0.39))
        else:  # drawn fallback: swap for the official logo when available
            self.draw_mark(s, 11.25, 0.52, 0.36, PURPLE); self.txt(s, "HCP Terraform", 11.65, 0.5, 1.7, 0.4, 14, FG, HEAD, wrap=False)
        self.counters.append(self.txt(s, "", 0.7, 7.05, 2, 0.3, 10, MUTE))
        if mark and self.mark: self.draw_mark(s, 9.9, 1.9, 3.2)
        return s

    def title_of(self, s, text):
        self.txt(s, text, 0.7, 0.8, 11, 0.9, 34, FG, HEAD)

    def notes(self, text):  # sidecar .md by default; embed_notes=True also writes real speaker notes (PowerPoint)
        self.notes_log.append((len(self.prs.slides), text))
        if self.embed_notes: self.prs.slides[-1].notes_slide.notes_text_frame.text = text

    # ---------- layouts
    def title(self, line1, line2, subtitle=""):
        s = self.new(mark=True)
        self.txt(s, line1, 0.9, 2.0, 8, 1.1, 66, PURPLE, HEAD); self.txt(s, line2, 0.9, 3.1, 8, 1.1, 66, FG, HEAD)
        if subtitle: self.txt(s, subtitle, 0.9, 4.6, 8, 0.8, 20, MUTE)

    def agenda(self, items):
        s = self.new(); self.title_of(s, "Agenda")
        for k, (n, name) in enumerate(items):
            y = 2.3 + k * 0.85
            self.rect(s, 0.7, y - 0.12, 11.9, 0.012, fill=LINE)
            self.txt(s, n, 0.7, y + 0.05, 1.2, 0.5, 16, PAL[k % 6], MONO)
            self.txt(s, name, 2.1, y - 0.02, 9.5, 0.6, 28, FG, HEAD)

    def section(self, num, name):
        s = self.new(mark=True)
        self.txt(s, num, 0.9, 2.5, 3, 0.5, 20, PURPLE, HEAD); self.txt(s, name, 0.9, 3.1, 8.0, 1.5, 54, FG, HEAD)

    def cards(self, ttl, cards):  # [(heading, [items])]; keep items to one line each
        s = self.new(); self.title_of(s, ttl); n = len(cards); w = (11.9 - 0.3 * (n - 1)) / n
        for j, (hd, items) in enumerate(cards):
            l = 0.7 + j * (w + 0.3); self.panel(s, l, 2.2, w, 4.3, PAL[j % 6])
            self.txt(s, hd, l + 0.3, 2.5, w - 0.4, 0.6, 18 if n > 3 else 22, FG, HEAD)
            for k, it in enumerate(items):
                self.txt(s, it, l + 0.3, 3.4 + k * 0.6, w - 0.4, 0.5, 12 if n > 3 else 15, MUTE)

    def rows(self, ttl, rows, callout=None):  # [(label, text)]
        s = self.new(); self.title_of(s, ttl)
        step = 0.78 if callout else (0.92 if len(rows) <= 5 else 0.66)
        for k, (a, b) in enumerate(rows):
            y = 2.3 + k * step
            self.rect(s, 0.7, y - 0.12, 11.9, 0.012, fill=LINE)
            self.txt(s, a, 0.7, y, 3.2, 0.5, 18, PAL[k % 6], HEAD); self.txt(s, b, 4.1, y + 0.03, 8.6, 0.5, 15, MUTE)
        if callout: self.txt(s, callout, 0.7, 6.3, 11.8, 0.5, 16, PURPLE, HEAD)

    def statement(self, lines):  # [(text, "mute"|"fg"|"acc")]
        s = self.new(mark=False); y = 1.9
        for text, role in lines:
            big = role != "acc"
            self.txt(s, text, 0.7, y, 12, 1.1, 36 if big else 22, {"mute": MUTE, "fg": FG, "acc": PURPLE}[role], HEAD)
            y += 0.95 if big else 0.6

    def big(self, ttl, tag, headline, sub):  # project slide: tag, one headline, 1-2 plain lines
        s = self.new(); self.title_of(s, ttl)
        if tag: self.tag(s, tag, 0.7, 1.9)
        self.txt(s, headline, 0.7, 2.9, 11.8, 2.2, 34, FG, HEAD)
        for n, ln in enumerate(sub if isinstance(sub, list) else [sub]):
            self.txt(s, ln, 0.7, 5.3 + n * 0.5, 11.8, 0.5, 18, FG if n == 0 and isinstance(sub, list) else MUTE)

    def cascade(self, ttl, headers, rows, callout):  # company goal -> priority -> work
        s = self.new(); self.title_of(s, ttl); xs = [0.7, 4.95, 9.2]; cols = [PAL[1], PAL[2], PAL[0]]
        for x, hd, c in zip(xs, headers, cols): self.txt(s, hd, x, 2.0, 3.45, 0.3, 11, c, MONO)
        for r, row in enumerate(rows):
            y = 2.45 + r * 1.25
            for c in range(3):
                self.panel(s, xs[c], y, 3.45, 1.05); self.rect(s, xs[c], y, 0.07, 1.05, fill=cols[c])
                self.txt(s, row[c * 2], xs[c] + 0.2, y + 0.14, 3.15, 0.4, 16, PURPLE if c == 2 else FG, HEAD)
                self.txt(s, row[c * 2 + 1], xs[c] + 0.2, y + 0.6, 3.15, 0.4, 12, MUTE)
                if c < 2: self.arrow(s, xs[c] + 3.85, y + 0.42)
        for n, ln in enumerate(callout):
            self.txt(s, ln, 0.7, 6.2 + n * 0.42, 12, 0.4, 15, PURPLE if n == 0 else FG, HEAD if n == 0 else BODY)

    def scorecard(self, ttl, tiles):  # [(category, big_number, plain_label)] x8, big stat tiles
        s = self.new(); self.title_of(s, ttl)
        for k, (cat, num, label) in enumerate(tiles):
            x, y, col = 0.7 + (k % 4) * 3.0, 1.95 + (k // 4) * 2.0, PAL[k % 6]
            self.panel(s, x, y, 2.8, 1.8, col)
            self.txt(s, cat, x + 0.2, y + 0.18, 2.5, 0.3, 11, MUTE, MONO)
            self.txt(s, num, x + 0.2, y + 0.5, 2.5, 0.7, 36, col, HEAD); self.txt(s, label, x + 0.2, y + 1.2, 2.45, 0.5, 12, FG)

    def shoutouts(self, ttl, rows):  # [(name, fun label, plain outcome)]
        s = self.new(); self.title_of(s, ttl)
        for k, (nm, label, what) in enumerate(rows):
            y = 2.0 + k * 0.6
            self.rect(s, 0.7, y - 0.1, 11.9, 0.012, fill=LINE)
            self.txt(s, nm, 0.7, y, 2.7, 0.4, 15, PAL[k % 6], HEAD); self.txt(s, label, 3.6, y + 0.01, 2.8, 0.4, 14, FG, HEAD)
            self.txt(s, what, 6.6, y + 0.03, 6.0, 0.4, 13, MUTE)

    def _circle(self, key, label, px=600):
        src = next((os.path.join(self.photos, key + e) for e in (".jpg", ".jpeg", ".png", ".webp")
                    if self.photos and os.path.exists(os.path.join(self.photos, key + e))), None)
        if src:
            im = Image.open(src).convert("RGB"); w, h = im.size; m = min(w, h)
            im = im.crop(((w - m) // 2, (h - m) // 2, (w + m) // 2, (h + m) // 2)).resize((px, px))
        else:  # initials placeholder until the real photo is dropped in
            im = Image.new("RGB", (px, px), (123, 66, 188)); d = ImageDraw.Draw(im)
            try: f = ImageFont.truetype("/System/Library/Fonts/Helvetica.ttc", px // 3)
            except Exception: f = ImageFont.load_default()
            d.text((px // 2, px // 2), "".join(p[0] for p in label.split()[:2]).upper(), fill="white", anchor="mm", font=f)
        mask = Image.new("L", (px, px), 0); ImageDraw.Draw(mask).ellipse((0, 0, px, px), fill=255); im.putalpha(mask)
        out = os.path.join(tempfile.gettempdir(), f"helios-photo-{key}.png"); im.save(out); return out

    def welcome(self, ttl, people):  # [(name, role, one-line blurb, photo_key)]
        s = self.new(); self.title_of(s, ttl)
        for k, (nm, role, blurb, key) in enumerate(people):
            x, col = 0.8 + k * 6.2, PAL[[0, 1, 2, 3][k % 4]]
            self.shape(s, MSO_SHAPE.OVAL, x - 0.07, 2.43, 2.74, 2.74, None, col, 2.5)
            s.shapes.add_picture(self._circle(key, nm), Inches(x), Inches(2.5), width=Inches(2.6), height=Inches(2.6))
            self.txt(s, nm, x + 3.0, 2.7, 3.2, 0.6, 24, FG, HEAD, wrap=False); self.txt(s, role, x + 3.0, 3.3, 3.0, 0.4, 17, col, HEAD)
            self.txt(s, blurb, x + 3.0, 3.85, 2.9, 1.2, 14, MUTE)

    def close(self, text, sub=""):
        s = self.new(mark=True); self.txt(s, text, 0.9, 2.2, 8.5, 2.4, 40, FG, HEAD)
        if sub: self.txt(s, sub, 0.9, 5.0, 8.5, 0.6, 18, PURPLE)

    # ---------- output
    def save(self, out):
        n = len(self.prs.slides)
        for i, tb in enumerate(self.counters, 1):
            r = tb.text_frame.paragraphs[0].add_run(); r.text = f"{i} / {n}"; r.font.size = Pt(10); r.font.name = BODY; r.font.color.rgb = rgb(MUTE)
        self.prs.save(out); self._fix_theme(out)
        if self.notes_log:
            open(out.replace(".pptx", "-notes.md"), "w").write("\n\n".join(f"Slide {i}: {t}" for i, t in self.notes_log))

    @staticmethod
    def _fix_theme(path):  # no Calibri/Arial/script fonts left in theme or masters
        tmp = path + ".tmp"
        with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
            for it in zin.infolist():
                data = zin.read(it.filename)
                if it.filename.startswith("ppt/theme/") and it.filename.endswith(".xml"):
                    x = data.decode("utf8")
                    x = re.sub(r'(<a:majorFont>\s*<a:latin typeface=")[^"]*(")', lambda m: m.group(1) + HEAD + m.group(2), x)
                    x = re.sub(r'(<a:minorFont>\s*<a:latin typeface=")[^"]*(")', lambda m: m.group(1) + BODY + m.group(2), x)
                    data = re.sub(r'<a:font script="[^"]*" typeface="[^"]*"/>', "", x).encode("utf8")
                elif it.filename.startswith(("ppt/slideMasters/", "ppt/slideLayouts/")) and it.filename.endswith(".xml"):
                    data = data.decode("utf8").replace('typeface="Arial"', f'typeface="{BODY}"').encode("utf8")
                zout.writestr(it, data)
        shutil.move(tmp, path)
